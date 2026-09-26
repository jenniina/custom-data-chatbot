from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import Mock

from docx import Document

from context_me.access import role_for
from context_me.chat import answer
from context_me.documents import chunks_for, convert_docx, dump_document, load_document
from context_me.library import Library


def sample_document():
    doc = Document()
    doc.add_heading("About me", level=1)
    doc.add_paragraph("My name is Elina. I live in Helsinki. Ääkköset work.")
    doc.add_heading("Interests", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Hobby"
    table.cell(0, 1).text = "Cycling"
    doc.add_paragraph("I cycle on weekends.")
    buffer = BytesIO()
    doc.save(buffer)
    return convert_docx(buffer.getvalue(), "profile.docx")


def fake_client():
    client = Mock()
    def embeddings(model, input):
        texts = [input] if isinstance(input, str) else input
        return SimpleNamespace(data=[SimpleNamespace(index=i, embedding=[1.0, 0.0] if "Helsinki" in t else [0.0, 1.0]) for i, t in enumerate(texts)])
    client.embeddings.create.side_effect = embeddings
    client.responses.create.return_value = SimpleNamespace(output_text="You live in Helsinki.")
    return client


class ConversionTests(unittest.TestCase):
    def test_round_trip_and_order(self):
        document = sample_document()
        self.assertEqual(load_document(dump_document(document)), document)
        text = "\n".join(c["text"] for c in document["chunks"])
        self.assertIn("Ääkköset", text)
        self.assertLess(text.index("Hobby | Cycling"), text.index("I cycle"))
        self.assertEqual(document["chunks"][1]["section"], "Interests")

    def test_chunk_boundaries(self):
        text = "x" * 12000
        parts = chunks_for(text)
        self.assertTrue(all(0 < len(p) <= 1800 for p in parts))
        rebuilt = parts[0] + "".join(p[200:] for p in parts[1:])
        self.assertEqual(rebuilt, text)

    def test_empty_doc_and_invalid_inputs(self):
        doc = Document()
        buffer = BytesIO()
        doc.save(buffer)
        with self.assertRaises(ValueError):
            convert_docx(buffer.getvalue(), "empty.docx")
        for raw in ("[]", "{}", "not json"):
            with self.assertRaises(ValueError):
                load_document(raw)
        with self.assertRaises(ValueError):
            convert_docx(b"bad", "bad.docx")

    def test_duplicate_chunk_rejected(self):
        document = sample_document()
        document["chunks"].append(document["chunks"][0])
        with self.assertRaises(ValueError):
            load_document(dump_document(document))


class SearchTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.library = Library(Path(self.temp.name) / "test.sqlite3")
        self.client = fake_client()
        self.document = sample_document()

    def tearDown(self):
        self.temp.cleanup()

    def test_search_persistence_dedup_and_removal(self):
        self.assertTrue(self.library.add(self.document, self.client, "embedding-model"))
        count = self.client.embeddings.create.call_count
        self.assertFalse(self.library.add(self.document, self.client, "embedding-model"))
        self.assertEqual(self.client.embeddings.create.call_count, count)
        reopened = Library(self.library.path)
        results = reopened.search("Helsinki", self.client)
        self.assertIn("Helsinki", results[0]["text"])
        self.assertEqual(results[0]["source"], "profile.docx")
        reopened.remove(self.document["document_id"])
        self.assertEqual(reopened.search("where", self.client), [])

    def test_failed_update_keeps_previous_data(self):
        self.library.add(self.document, self.client, "embedding-model")
        self.document["chunks"][0]["text"] = "Updated data"
        self.client.embeddings.create.side_effect = RuntimeError("API offline")
        with self.assertRaises(RuntimeError):
            self.library.add(self.document, self.client, "embedding-model")
        results = self.library.search("Helsinki", fake_client())
        self.assertIn("Helsinki", results[0]["text"])

    def test_wrong_model_rejected(self):
        self.library.add(self.document, self.client, "embedding-model")
        with self.assertRaises(ValueError):
            self.library.add(self.document, self.client, "other-model")

    def test_grounded_response_and_history(self):
        self.library.add(self.document, self.client, "embedding-model")
        history = [{"role": "user", "content": "Where do I live?"}, {"role": "assistant", "content": "Helsinki"}]
        text, sources = answer("What about Helsinki?", history, self.library, self.client, "chat-model")
        self.assertEqual(text, "You live in Helsinki.")
        args = self.client.responses.create.call_args.kwargs
        self.assertFalse(args["store"])
        self.assertIn("Retrieved sources", args["input"][-1]["content"])
        self.assertIn("Where do I live?", self.client.embeddings.create.call_args.kwargs["input"])
        self.assertEqual(history[0]["content"], "Where do I live?")
        self.assertTrue(sources)

    def test_empty_library_does_not_call_api(self):
        with self.assertRaises(ValueError):
            answer("Where do I live?", [], self.library, self.client, "model")
        self.client.responses.create.assert_not_called()
        self.client.embeddings.create.assert_not_called()


class AccessTests(unittest.TestCase):
    def test_roles_and_wrong_password(self):
        self.assertEqual(role_for("admin", "admin"), "admin")
        self.assertIsNone(role_for("wrong", "admin"))
        self.assertIsNone(role_for("", ""))


if __name__ == "__main__":
    unittest.main()
