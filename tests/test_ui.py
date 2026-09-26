from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from docx import Document
from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import Client, TestCase, override_settings
from context_me.documents import dump_document
from context_me.library import Library
from tests.test_core import fake_client, sample_document

PASSWORD = "a-long-test-admin-password"

class InterfaceTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = get_user_model().objects.create_user("admin", password=PASSWORD, is_staff=True)

    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        override = override_settings(DATA_DIR=Path(self.tmp.name))
        override.enable()
        self.addCleanup(override.disable)
        cache.clear()
        self.library = Library(Path(self.tmp.name) / "library.sqlite3")
        self.api = fake_client()
        patcher = patch("webapp.views.api_client", return_value=self.api)
        patcher.start()
        self.addCleanup(patcher.stop)

    def login(self):
        self.client.force_login(self.admin)

    def upload(self, plain=False):
        word = Document()
        word.add_paragraph("I live in Helsinki." if plain else "EN: I live in Helsinki.")
        if not plain:
            word.add_paragraph("FI: Asun Helsingissä.")
        output = BytesIO()
        word.save(output)
        return self.client.post("/admin/convert/", {"action": "upload", "document": SimpleUploadedFile("profile.docx", output.getvalue())})

    def test_public_reader_has_no_admin_controls(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, 'type="file"')
        self.assertNotContains(response, "API key")
        self.assertNotContains(response, "Retrieved sources")
        self.assertContains(response, 'id="main"')
        self.assertNotContains(response, "MutationObserver")

    def test_only_assistant_replies_render_markdown_without_numbering_messages(self):
        session = self.client.session
        session["conversation"] = [
            {"role": "user", "content": "**My question** <script>bad</script>"},
            {"role": "assistant", "content": "A **designer**. [1]\n\n[Portfolio](https://example.com)"},
        ]
        session.save()
        response = self.client.get("/")
        self.assertContains(response, "<strong>designer</strong>")
        self.assertContains(response, '<a href="https://example.com">Portfolio</a>')
        self.assertContains(response, "**My question** &lt;script&gt;bad&lt;/script&gt;")
        self.assertNotContains(response, '<ol class="conversation">')

    def test_all_admin_routes_require_login(self):
        for path in ["/admin/library/", "/admin/convert/", "/admin/appearance/", "/admin/preview/", "/admin/onedrive/",
                     "/admin/remove/anything/", "/admin/download/anything/"]:
            for method in (self.client.get, self.client.post):
                self.assertEqual(method(path).status_code, 302, path)

    def test_nonstaff_user_is_not_administrator(self):
        user = get_user_model().objects.create_user("reader", password=PASSWORD)
        self.client.force_login(user)
        self.assertEqual(self.client.get("/admin/library/").status_code, 302)
        self.client.logout()
        response = self.client.post("/admin/login/", {"username": "reader", "password": PASSWORD})
        self.assertContains(response, "does not have administrator access")

    def test_login_logout_and_password_rotation(self):
        response = self.client.post("/admin/login/", {"username": "admin", "password": PASSWORD})
        self.assertRedirects(response, "/admin/library/")
        self.assertEqual(self.client.get("/admin/logout/").status_code, 405)
        self.admin.set_password(PASSWORD + "-changed")
        self.admin.save()
        self.assertEqual(self.client.get("/admin/library/").status_code, 302)
        self.login()
        session = self.client.session
        session["drive_cache"] = "sensitive-token"
        session["preview"] = sample_document()
        session.save()
        self.client.post("/admin/logout/")
        self.assertNotIn("drive_cache", self.client.session)
        self.assertNotIn("preview", self.client.session)

    def test_login_rate_limit(self):
        for _ in range(11):
            response = self.client.post("/admin/login/", {"username": "admin", "password": "wrong"})
        self.assertContains(response, "Too many sign-in attempts")

    def test_csrf_and_safe_methods(self):
        browser = Client(enforce_csrf_checks=True)
        browser.force_login(self.admin)
        for path in ["/", "/chat/clear/", "/admin/appearance/", "/admin/convert/", "/admin/logout/"]:
            self.assertEqual(browser.post(path, {}).status_code, 403)
        self.assertEqual(browser.get("/chat/clear/").status_code, 405)

    def test_public_chat_sources_never_render_and_sessions_are_separate(self):
        self.library.add(sample_document(), self.api, "embedding-model")
        response = self.client.post("/", {"question": "Where do I live?", "language": "en"}, follow=True)
        self.assertContains(response, "You live in Helsinki.")
        self.assertNotContains(response, "Retrieved sources")
        self.assertEqual(self.client.session["conversation"][-1]["sources"], [])
        self.assertNotContains(Client().get("/"), "You live in Helsinki.")
        session = self.client.session
        history = session["conversation"]
        history[-1]["sources"] = [{"source": "private-name", "text": "RAW_SOURCE_SENTINEL"}]
        session["conversation"] = history
        session.save()
        self.assertNotContains(self.client.get("/"), "RAW_SOURCE_SENTINEL")
        self.assertNotIn("Helsinki", self.client.cookies["sessionid"].value)

    def test_question_failure_retains_input_and_no_history(self):
        self.library.add(sample_document(), self.api, "embedding-model")
        self.api.responses.create.side_effect = RuntimeError("secret-api-detail")
        response = self.client.post("/", {"question": "Where do I live?", "language": "fi"})
        self.assertContains(response, "temporarily unavailable")
        self.assertContains(response, "Where do I live?")
        self.assertNotContains(response, "secret-api-detail")
        self.assertNotIn("conversation", self.client.session)

    def test_clear_conversation(self):
        session = self.client.session
        session["conversation"] = [{"role": "user", "content": "Private question"}]
        session.save()
        self.client.post("/chat/clear/")
        self.assertNotIn("conversation", self.client.session)

    def test_bilingual_conversion_publishes_directly_and_deduplicates(self):
        self.login()
        self.assertRedirects(self.upload(), "/admin/preview/")
        self.api.embeddings.create.assert_not_called()
        document = self.client.session["preview"]
        self.assertTrue(any(r["fi"] for r in document["records"]))
        token = self.client.session["preview_id"]
        response = self.client.post("/admin/preview/", {"preview_id": token}, follow=True)
        self.assertContains(response, "Document added")
        self.assertEqual(len(self.library.documents()), 1)
        count = self.api.embeddings.create.call_count
        self.client.post("/admin/preview/", {"preview_id": token})
        self.assertEqual(self.api.embeddings.create.call_count, count)
        download = self.client.get("/admin/download/" + token + "/")
        self.assertEqual(download.status_code, 200)
        self.assertIn(b'"schema_version": 2', b"".join(download.streaming_content))

    def test_plain_single_language_and_json_upload(self):
        self.login()
        self.upload(plain=True)
        document = self.client.session["preview"]
        self.assertIn("Helsinki", document["chunks"][0]["text"])
        response = self.client.post("/admin/convert/", {"action": "upload", "document":
            SimpleUploadedFile("profile.json", dump_document(document).encode())})
        self.assertRedirects(response, "/admin/preview/")

    def test_invalid_upload_and_stale_preview(self):
        self.login()
        response = self.client.post("/admin/convert/", {"action": "upload", "document": SimpleUploadedFile("bad.docx", b"not a zip")})
        self.assertContains(response, "not a readable Word")
        self.upload()
        self.assertEqual(self.client.post("/admin/preview/", {"preview_id": "stale"}).status_code, 404)
        self.assertEqual(self.client.get("/admin/download/stale/").status_code, 404)
        self.api.embeddings.create.assert_not_called()

    def test_failed_indexing_keeps_library_unchanged(self):
        self.login()
        self.upload()
        self.api.embeddings.create.side_effect = RuntimeError("secret-api-detail")
        response = self.client.post("/admin/preview/", {"preview_id": self.client.session["preview_id"]})
        self.assertContains(response, "could not be completed")
        self.assertNotContains(response, "secret-api-detail")
        self.assertEqual(self.library.documents(), [])

    def test_preview_is_private_and_html_escaped(self):
        self.login()
        document = sample_document()
        document["chunks"][0]["text"] = '<script>alert("x")</script>'
        session = self.client.session
        session["preview"] = document
        session["preview_id"] = "test-token"
        session.save()
        response = self.client.get("/admin/preview/")
        self.assertContains(response, "&lt;script&gt;")
        self.assertNotContains(response, '<script>alert("x")</script>')
        other = Client()
        other.force_login(self.admin)
        self.assertEqual(other.get("/admin/download/test-token/").status_code, 404)

    def test_removal_requires_post_confirmation(self):
        document = sample_document()
        self.library.add(document, self.api, "embedding-model")
        self.login()
        path = "/admin/remove/" + document["document_id"] + "/"
        self.assertContains(self.client.get(path), "Cancel removal")
        self.assertEqual(len(self.library.documents()), 1)
        self.client.post(path)
        self.assertEqual(self.library.documents(), [])

    def test_appearance_persists_and_is_escaped(self):
        self.login()
        self.client.post("/admin/appearance/", {"title": "Ask Jenniina", "caption": "<script>bad</script>", "explanation": "Kysy suomeksi"})
        response = Client().get("/")
        self.assertContains(response, "Ask Jenniina")
        self.assertContains(response, "Kysy suomeksi")
        self.assertNotContains(response, "<script>bad</script>")
        self.assertEqual(self.library.presentation()["title"], "Ask Jenniina")

    @patch("webapp.views.read_public_document")
    def test_public_onedrive_review_before_publish(self, read):
        self.login()
        read.return_value = sample_document()
        response = self.client.post("/admin/convert/", {"action": "link", "url": "https://1drv.ms/example", "source_name": ""})
        self.assertRedirects(response, "/admin/preview/")
        self.api.embeddings.create.assert_not_called()
        read.assert_called_once_with(url="https://1drv.ms/example", source_name="")

    @patch("webapp.views.OneDrive")
    def test_private_onedrive_session_and_disconnect(self, drive):
        self.login()
        drive.return_value.begin_sign_in.return_value = {"user_code": "ABCDEF", "device_code": "private-device-code", "interval": 5}
        drive.return_value.serialize_cache.return_value = "secret-cache"
        self.client.post("/admin/onedrive/", {"action": "connect"})
        self.assertEqual(self.client.session["drive_cache"], "secret-cache")
        self.assertNotIn("secret-cache", self.client.cookies["sessionid"].value)
        drive.return_value.finish_sign_in.return_value = True
        self.client.post("/admin/onedrive/", {"action": "poll"})
        self.assertTrue(self.client.session["drive_ready"])
        drive.return_value.read_document.return_value = sample_document()
        self.assertRedirects(self.client.post("/admin/onedrive/", {"action": "load", "path": "profile.docx"}), "/admin/preview/")
        self.client.post("/admin/onedrive/", {"action": "disconnect"})
        self.assertNotIn("drive_cache", self.client.session)
        self.assertNotIn("preview", self.client.session)

    def test_security_response_headers(self):
        response = self.client.get("/")
        self.assertIn("no-store", response["Cache-Control"])
        self.assertEqual(response["X-Frame-Options"], "DENY")
        self.assertNotIn("unsafe-inline", response["Content-Security-Policy"])
