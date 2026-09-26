"""Django browser audit using temporary data, mocked OpenAI and local axe-core.

Install requirements-dev.txt and place axe.min.js in data/accessibility/.
"""
import json
import os
from pathlib import Path
import socket
import subprocess
import sys
import tempfile
import time
from urllib.request import urlopen

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ["CONTEXT_ME_LOAD_LOCAL_CONFIG"] = "false"
from playwright.sync_api import sync_playwright, expect
from context_me.library import Library
from tests.test_core import fake_client, sample_document

def main():
    output = ROOT / "data/accessibility"
    output.mkdir(parents=True, exist_ok=True)
    axe = output / "axe.min.js"
    if not axe.exists():
        raise SystemExit("Place axe-core axe.min.js in data/accessibility/ before running this audit.")
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp)
        library = Library(path / "library.sqlite3")
        library.add(sample_document(), fake_client(), "embedding-model")
        wrapper = path / "audit_app.py"
        wrapper.write_text(
            f"import sys\nsys.path.insert(0, {str(ROOT)!r})\n"
            "import threading, time, os\nfrom pathlib import Path\n"
            f"def stop_when_requested():\n    while not Path({str(path / 'stop')!r}).exists(): time.sleep(.1)\n    os._exit(0)\n"
            "threading.Thread(target=stop_when_requested, daemon=True).start()\n"
            "import django\ndjango.setup()\nfrom unittest.mock import patch\n"
            "from tests.test_core import fake_client\nfrom serve import main\n"
            "client = fake_client()\n"
            "client.responses.create.return_value.output_text = 'You live in Helsinki. **A designer.** [1]\\n\\n[Portfolio](https://example.com/portfolio)'\n"
            "with patch('webapp.views.api_client', return_value=client):\n    main()\n", encoding="utf-8")
        with socket.socket() as sock:
            sock.bind(("127.0.0.1", 0))
            port = sock.getsockname()[1]
        env = dict(os.environ, CONTEXT_ME_LOAD_LOCAL_CONFIG="false", CONTEXT_ME_HOSTED="false",
                   DATABASE_URL="", ADMIN_USERNAME="audit-editor", ADMIN_PASSWORD="accessibility-test-password", DJANGO_SETTINGS_MODULE="webapp.settings",
                   DJANGO_ALLOWED_HOSTS="localhost,127.0.0.1", OPENAI_API_KEY="fake-audit-key",
                   CONTEXT_ME_DATA_DIR=str(path), ONEDRIVE_CLIENT_ID="", OPENAI_EMBEDDING_MODEL="embedding-model",
                   PORT=str(port))
        env.pop("K_SERVICE", None)
        with (output / "django-server.log").open("w") as log:
            server = subprocess.Popen([sys.executable, str(wrapper)], cwd=ROOT, env=env, stdout=log, stderr=log)
            try:
                url = f"http://127.0.0.1:{port}"
                for _ in range(60):
                    if server.poll() is not None:
                        raise RuntimeError("Audit server exited. See django-server.log.")
                    try:
                        urlopen(url, timeout=1).close()
                        break
                    except OSError:
                        time.sleep(.5)
                with sync_playwright() as p:
                    browser = p.chromium.launch(channel="msedge", headless=True)
                    context = browser.new_context(viewport={"width": 1280, "height": 900})
                    # Browser instrumentation; no changes to the app's CSP or markup.
                    context.add_init_script(path=str(axe))
                    page = context.new_page()
                    results = []
                    def audit(name):
                        report = page.evaluate("""async () => {
                            const r = await axe.run(document, {runOnly: {type:'tag', values:['wcag2a','wcag2aa','wcag21a','wcag21aa','wcag22aa','best-practice']}});
                            return {violations:r.violations.map(v=>({id:v.id,nodes:v.nodes.map(n=>({target:n.target,summary:n.failureSummary}))})),
                                    incomplete:r.incomplete.map(v=>({id:v.id,nodes:v.nodes.map(n=>({target:n.target,summary:n.failureSummary,checks:n.any}))}))};
                        }""")
                        report.update(state=name, horizontal_overflow=page.evaluate("document.documentElement.scrollWidth > innerWidth"))
                        # Supplemental checks for native textarea contrast that axe cannot resolve.
                        report["textarea_checks"] = page.locator("textarea").evaluate_all("""elements => {
                          function lum(rgb) {
                            const c = rgb.match(/[0-9.]+/g).slice(0,3).map(v=>Number(v)/255).map(v=>v<=.04045 ? v/12.92 : ((v+.055)/1.055)**2.4);
                            return c[0]*.2126+c[1]*.7152+c[2]*.0722;
                          }
                          return elements.map(e=>{
                            e.scrollIntoView({block:'center'});
                            const s=getComputedStyle(e), a=lum(s.color), b=lum(s.backgroundColor), r=e.getBoundingClientRect();
                            const unobscured=[[.2,.2],[.8,.2],[.5,.5],[.2,.8],[.8,.8]].every(([x,y])=>document.elementFromPoint(r.left+r.width*x,r.top+r.height*y)===e);
                            return {id:e.id,foreground:s.color,background:s.backgroundColor,contrast:(Math.max(a,b)+.05)/(Math.min(a,b)+.05),unobscured};
                          });
                        }""")
                        assert all(c["contrast"] >= 4.5 and c["unobscured"] for c in report["textarea_checks"]), report["textarea_checks"]
                        results.append(report)
                        page.screenshot(path=str(output / ("django-" + name + ".png")), full_page=True)
                        (output / "django-report.json").write_text(json.dumps(results, indent=2), encoding="utf-8")
                    page.goto(url)
                    expect(page.get_by_role("heading", name="Context Me", exact=True)).to_be_visible()
                    assert not page.locator("details.menu").evaluate("(e)=>e.open")
                    page.keyboard.press("Tab")
                    expect(page.get_by_role("link", name="Skip to main content")).to_be_focused()
                    page.keyboard.press("Enter")
                    expect(page.locator("#main")).to_be_focused()
                    for _ in range(10):
                        page.keyboard.press("Tab")
                        if page.locator("#id_question").evaluate("(e)=>e===document.activeElement"):
                            break
                    expect(page.locator("#id_question")).to_be_focused()
                    assert page.locator("#id_question").evaluate("(e)=>getComputedStyle(e).outlineStyle") != "none"
                    audit("public-desktop")
                    page.get_by_label("What would you like to know?").fill("Where do I live?")
                    page.get_by_label("What would you like to know?").press("Enter")
                    expect(page.locator("#latest-answer")).to_contain_text("You live in Helsinki.")
                    expect(page.locator("#latest-answer")).to_be_focused()
                    assert "Retrieved sources" not in page.content()
                    expect(page.locator("#latest-answer strong")).to_have_text("A designer.")
                    expect(page.locator("#latest-answer a")).to_have_attribute("href", "https://example.com/portfolio")
                    audit("public-answer")
                    page.set_viewport_size({"width": 320, "height": 800})
                    audit("public-320px")
                    page.set_viewport_size({"width": 1280, "height": 900})
                    page.locator("summary").focus()
                    page.keyboard.press("Enter")
                    assert page.locator("details.menu").evaluate("(e)=>e.open")
                    page.get_by_role("link", name="Administrator sign-in").click()
                    audit("login")
                    page.get_by_label("Username").fill("audit-editor")
                    page.get_by_label("Password").fill("wrong")
                    page.get_by_role("button", name="Sign in", exact=True).click()
                    expect(page.locator("#form-errors")).to_be_focused()
                    audit("login-error")
                    page.get_by_label("Password").fill("accessibility-test-password")
                    page.get_by_role("button", name="Sign in", exact=True).click()
                    expect(page.get_by_role("heading", name="My library", exact=True)).to_be_visible()
                    audit("admin-library")
                    page.get_by_role("link", name="Remove profile.docx").click()
                    audit("remove-confirmation")
                    page.get_by_role("link", name="Cancel removal").focus()
                    page.keyboard.press("Enter")
                    assert len(library.documents()) == 1
                    page.get_by_role("link", name="Add a document").click()
                    audit("converter")
                    from docx import Document
                    fixture = path / "accessible-upload.docx"
                    word = Document()
                    word.add_paragraph("EN: I live in Helsinki.")
                    word.add_paragraph("FI: Asun Helsingissä.")
                    word.save(fixture)
                    page.get_by_label("Word document or converted JSON").set_input_files(str(fixture))
                    page.get_by_role("button", name="Review document", exact=True).click()
                    expect(page.get_by_role("heading", name="Review document", exact=True)).to_be_visible()
                    page.get_by_text("Review separated languages", exact=True).click()
                    audit("converter-preview")
                    page.set_viewport_size({"width": 320, "height": 800})
                    audit("preview-320px")
                    page.set_viewport_size({"width": 1280, "height": 900})
                    page.get_by_role("button", name="Add to chatbot library").click()
                    expect(page.get_by_text("Document added to the chatbot library.")).to_be_visible()
                    assert len(library.documents()) == 2
                    audit("publish-success")
                    page.goto(url + "/admin/appearance/")
                    audit("appearance")
                    page.get_by_label("Title").fill("Ask about me")
                    page.get_by_role("button", name="Save appearance").click()
                    expect(page.get_by_text("Appearance saved.")).to_be_visible()
                    page.goto(url + "/admin/onedrive/")
                    audit("private-onedrive-unconfigured")
                    # The central workflow works without JavaScript as well.
                    nojs = browser.new_context(java_script_enabled=False)
                    plain = nojs.new_page()
                    plain.goto(url)
                    plain.get_by_label("What would you like to know?").fill("Where do I live?")
                    plain.get_by_role("button", name="Send question").click()
                    expect(plain.locator("#latest-answer")).to_contain_text("Helsinki")
                    nojs.close()
                    print(json.dumps(results, indent=2))
                    assert not any(r["violations"] or r["horizontal_overflow"] for r in results), "See django-report.json"
                    browser.close()
            finally:
                (path / "stop").write_text("stop", encoding="utf-8")
                server.wait(timeout=15)

if __name__ == "__main__":
    main()
