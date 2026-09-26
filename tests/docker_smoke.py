"""Smoke-test the built image with disposable data and credentials; no API calls.

Run: python tests/docker_smoke.py [image-tag] [container-port]
Requires Docker and the project's local Python dependencies.
"""
import os
from io import BytesIO
import re
import secrets
import subprocess
import sys
import time

from docx import Document
import requests


def docker(*args, **kwargs):
    return subprocess.check_output(["docker", *args], text=True, **kwargs).strip()


def main():
    image = sys.argv[1] if len(sys.argv) > 1 else "context-me:django-test"
    container_port = int(sys.argv[2]) if len(sys.argv) > 2 else 8080
    env = dict(os.environ, ADMIN_USERNAME="smoke-" + secrets.token_hex(8), ADMIN_PASSWORD=secrets.token_urlsafe(24),
               DJANGO_SECRET_KEY=secrets.token_urlsafe(64))
    container = docker("run", "-d", "--rm", "-p", f"127.0.0.1::{container_port}",
                       "-e", f"PORT={container_port}", "-e", "K_SERVICE=context-me-smoke",
                       "-e", "ADMIN_USERNAME", "-e", "ADMIN_PASSWORD", "-e", "DJANGO_SECRET_KEY",
                       "-e", "DJANGO_ALLOWED_HOSTS=127.0.0.1,localhost",
                       "-e", "ALLOW_TEMPORARY_DATABASE=true",
                       "-e", "CONTEXT_ME_LOAD_LOCAL_CONFIG=false", image, env=env)
    try:
        port = docker("port", container, f"{container_port}/tcp").rsplit(":", 1)[1]
        base = "http://127.0.0.1:" + port
        session = requests.Session()
        session.headers["X-Forwarded-Proto"] = "https"

        def request(method, path, **kwargs):
            # Emulate an HTTPS terminating proxy over a loopback-only test port.
            # Real browsers receive these Secure cookies over actual HTTPS.
            headers = {"Cookie": "; ".join(k + "=" + v for k, v in session.cookies.get_dict().items()),
                       "Origin": base.replace("http:", "https:")}
            return session.request(method, base + path, headers=headers, timeout=15,
                                   allow_redirects=False, **kwargs)

        for _ in range(60):
            try:
                response = request("GET", "/")
                if response.status_code == 200:
                    break
            except requests.ConnectionError:
                pass
            time.sleep(.5)
        else:
            raise AssertionError("Container did not become ready.")

        assert "Ask a question" in response.text
        assert "Retrieved sources" not in response.text
        assert 'type="file"' not in response.text
        assert "no-store" in response.headers["Cache-Control"]
        css = re.search(r'href="(/static/[^\"]+\.css)"', response.text).group(1)
        assert request("GET", css).status_code == 200
        assert request("GET", "/admin/library/").status_code == 302
        assert request("POST", "/chat/clear/").status_code == 403
        assert requests.get(base, timeout=10, allow_redirects=False).status_code == 301

        request("GET", "/admin/login/")
        response = request("POST", "/admin/login/", data={
            "csrfmiddlewaretoken": session.cookies["csrftoken"],
            "username": env["ADMIN_USERNAME"], "password": env["ADMIN_PASSWORD"]})
        assert response.status_code == 302, "Administrator login failed."
        assert session.cookies["sessionid"]
        assert all(cookie.secure for cookie in session.cookies)
        assert request("GET", "/admin/library/").status_code == 200

        doc = Document()
        doc.add_paragraph("EN: I live in Helsinki.")
        doc.add_paragraph("FI: Asun Helsingissä.")
        raw = BytesIO()
        doc.save(raw)
        response = request("POST", "/admin/convert/", data={
            "action": "upload", "csrfmiddlewaretoken": session.cookies["csrftoken"]},
            files={"document": ("test-profile.docx", raw.getvalue())})
        assert response.status_code == 302
        preview = request("GET", "/admin/preview/")
        assert "Asun Helsingissä" in preview.text
        url = re.search(r'href="(/admin/download/[^\"]+)"', preview.text).group(1)
        document = request("GET", url).json()
        assert document["schema_version"] == 2
        assert any(record["fi"] for record in document["records"])

        output = docker("exec", container, "python", "manage.py", "check", "--deploy")
        assert "no issues" in output
        inspection = docker("exec", container, "python", "-c",
            "import os,json; from pathlib import Path; "
            "p=Path('/app'); "
            "assert os.getuid()!=0; "
            "assert not any(x.suffix in ('.docx','.json','.sqlite3','.db','.pem','.key') "
            "or x.name in ('.env','secrets.toml','.django-secret') for x in p.rglob('*') if x.is_file()); "
            "assert not (p/'.git').exists(); assert not (p/'.venv').exists(); "
            "assert (p/'webapp/management/commands/sync_admin.py').exists(); "
            "print('Image runs as non-root and contains no personal-data or secret files.')")
        print(inspection)
        print("Docker smoke checks passed: startup, public chat page, static files, HTTPS handling, "
              "CSRF, admin login, DOCX preview, bilingual JSON download, deployment checks.")
    finally:
        subprocess.run(["docker", "rm", "-f", container], check=True, stdout=subprocess.DEVNULL)


if __name__ == "__main__":
    main()
