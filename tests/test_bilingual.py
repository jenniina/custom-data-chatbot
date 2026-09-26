from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from docx import Document

from context_me.documents import convert_docx, dump_document, load_document
from context_me.library import Library
from tests.test_core import fake_client


def convert_lines(lines):
    doc = Document()
    for line in lines:
        if isinstance(line, tuple):
            doc.add_heading(line[0], level=1)
        else:
            doc.add_paragraph(line)
    raw = BytesIO()
    doc.save(raw)
    return convert_docx(raw.getvalue(), "bilingual.docx")


class BilingualTests(unittest.TestCase):
    def test_adjacent_and_inline_pairs(self):
        document = convert_lines([("Profile",), "EN: I live in Helsinki.", "FI: Asun Helsingissä.",
                                  "[service.design] EN: Web design | FI: Verkkosuunnittelu"])
        pair, inline = document["records"][1:]
        self.assertEqual(pair["en"], "I live in Helsinki.")
        self.assertEqual(pair["fi"], "Asun Helsingissä.")
        self.assertEqual(pair["original_text"], "EN: I live in Helsinki.\n\nFI: Asun Helsingissä.")
        self.assertEqual(inline["key"], "service.design")
        self.assertEqual(inline["fi"], "Verkkosuunnittelu")
        self.assertEqual(load_document(dump_document(document)), document)

    def test_shared_values_and_missing_translations(self):
        document = convert_lines(["[person.name] Full name / Koko nimi: Elina",
                                  "[site.url] https://example.invalid/fi/en",
                                  "EN: English only", "[other.fact] FI: Vain suomeksi"])
        records = document["records"]
        self.assertEqual(records[0]["shared"], "Full name / Koko nimi: Elina")
        self.assertEqual(records[0]["en"], "")
        self.assertEqual(records[1]["shared"], "https://example.invalid/fi/en")
        self.assertEqual(records[2]["fi"], "")
        self.assertEqual(records[3]["en"], "")

    def test_pairing_does_not_cross_sections_or_shared_text(self):
        document = convert_lines([("One",), "EN: First", ("Two",), "FI: Toinen",
                                  "A shared fact", "EN: Another"])
        self.assertFalse(any(r["en"] and r["fi"] for r in document["records"]))

    def test_reversed_and_keyed_pairs(self):
        document = convert_lines(["[fact.a] FI: Suomi", "[fact.a] EN: English"])
        self.assertEqual(len(document["records"]), 1)
        self.assertEqual(document["records"][0]["en"], "English")

    def test_ambiguous_labels_not_discarded(self):
        text = "EN: First | EN: Second | FI: Toinen"
        record = convert_lines([text])["records"][0]
        self.assertEqual(record["shared"], text)
        self.assertEqual(record["en"], "")

    def test_long_pair_keeps_complete_record(self):
        en, fi = "English " * 700, "Suomeksi " * 700
        document = convert_lines(["EN: " + en, "FI: " + fi])
        record = document["records"][0]
        self.assertEqual(record["en"], en.strip())
        self.assertEqual(record["fi"], fi.strip())
        self.assertTrue(all(len(c["text"]) <= 4000 for c in document["chunks"]))
        self.assertTrue(all(c["record_ids"] == [record["id"]] for c in document["chunks"]))
        self.assertEqual(load_document(dump_document(document)), document)

    def test_ordinary_pair_not_split_at_packing_target(self):
        document = convert_lines(["EN: " + "e" * 1400, "FI: " + "f" * 1400])
        self.assertEqual(len(document["chunks"]), 1)

    def test_v1_compatibility_and_v2_validation(self):
        document = convert_lines(["EN: Hello", "FI: Hei"])
        legacy = {k: v for k, v in document.items() if k != "records"}
        legacy["schema_version"] = 1
        self.assertEqual(load_document(dump_document(legacy)), legacy)
        document["records"][0]["fi"] = "Changed"
        with self.assertRaisesRegex(ValueError, "do not match"):
            load_document(dump_document(document))

    def test_both_languages_indexed_and_v1_replaced(self):
        document = convert_lines(["EN: I live in Helsinki.", "FI: Asun Helsingissä."])
        legacy = {k: v for k, v in document.items() if k != "records"}
        legacy["schema_version"] = 1
        client = fake_client()
        with TemporaryDirectory() as tmp:
            library = Library(Path(tmp) / "library.sqlite3")
            library.add(legacy, client, "embedding-model")
            self.assertTrue(library.add(document, client, "embedding-model"))
            self.assertEqual(len(library.documents()), 1)
            indexed = client.embeddings.create.call_args.kwargs["input"][0]
            self.assertIn("I live in Helsinki", indexed)
            self.assertIn("Asun Helsingissä", indexed)
            self.assertFalse(library.add(document, client, "embedding-model"))


if __name__ == "__main__":
    unittest.main()
